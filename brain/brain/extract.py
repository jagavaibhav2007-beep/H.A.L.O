"""Layer 0 of the document-ingestion design (systemdesign/13-document-ingestion.md):
an extension -> markdown dispatcher. Deterministic, offline, no LLM, no network.

Mirrors microsoft/markitdown's dispatch shape (MIT) but swaps its pdfminer PDF
path for pypdfium2 + a pypdf fallback -- see the design doc's Provenance table
for why (AGPL/quality tradeoffs of the alternatives).

`extract_text(path)` is the one public entry point every caller (file_read,
later doc_digest) should route through.
"""

from __future__ import annotations

from pathlib import Path

import mammoth
import openpyxl
import pypdfium2 as pdfium
from markdownify import markdownify
from pypdf import PdfReader
from brain.extract_worker import PDF_OUTPUT_BYTES, PDF_PAGE_CAP

# ponytail: per-sheet row cap keeps a whole workbook cheap to page through the
# chat loop; raise it (or teach doc_digest to page sheets) if a real workflow
# needs full sheets returned in one call.
_XLSX_ROW_CAP = 200
# PDF parsing is worker-only with OS memory/time limits as well as these caps.
_XLSX_SHEET_CAP = 20
_PDF_PAGE_CAP = PDF_PAGE_CAP
# Hard refusal above this. Every converter below reads the whole file into
# memory, and file_read's caller is Tier 1 inside roots (no approval), so an
# unbounded read is an unattended OOM of the Brain.
_MAX_BYTES = 64 * 1024 * 1024


def _truncation_note(kind: str, cap: int, total: int) -> str:
    return f"\n\n... [truncated at {cap} {kind} of {total} total]"


class _OutputLimit(ValueError):
    pass


def _extract_pdf(path: Path) -> str:
    """Worker-only parser. Call extract_text for the contained public API."""
    text = ""
    total = 0
    try:
        pdf = pdfium.PdfDocument(str(path))
        try:
            total = len(pdf)
            parts = []
            output_bytes = 0
            for i in range(min(total, _PDF_PAGE_CAP)):
                page = pdf[i]
                try:
                    textpage = page.get_textpage()
                    try:
                        if textpage.count_chars() > PDF_OUTPUT_BYTES:
                            raise _OutputLimit("PDF output limit exceeded")
                        part = textpage.get_text_range()
                    finally:
                        textpage.close()
                finally:
                    page.close()
                output_bytes += len(part.encode("utf-8")) + 2
                if output_bytes > PDF_OUTPUT_BYTES - 128:
                    raise _OutputLimit("PDF output limit exceeded")
                parts.append(part)
        finally:
            pdf.close()
        text = "\n\n".join(parts)
    except _OutputLimit:
        raise
    except Exception:
        text = ""  # fall through to the pypdf fallback below
    if not text.strip():
        try:
            reader = PdfReader(str(path))
            if reader.is_encrypted:
                # Distinct from a scanned page: the text may be right there,
                # locked. Say so, so the remedy (supply the password) is honest.
                raise ValueError(f"{path.name} is an encrypted/password-protected PDF")
            total = len(reader.pages)
            parts = []
            output_bytes = 0
            for i in range(min(total, _PDF_PAGE_CAP)):
                part = reader.pages[i].extract_text() or ""
                output_bytes += len(part.encode("utf-8")) + 2
                if output_bytes > PDF_OUTPUT_BYTES - 128:
                    raise ValueError("PDF output limit exceeded")
                parts.append(part)
            text = "\n\n".join(parts)
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError(f"could not parse PDF {path.name}: {exc}") from exc
    if text.strip() and total > _PDF_PAGE_CAP:
        text += _truncation_note("pages", _PDF_PAGE_CAP, total)
    if not text.strip():
        raise ValueError(
            f"no extractable text in {path.name} (scanned PDF? Layer 0 is text-only, no OCR yet)"
        )
    return text


def _pdf_pages(path: Path) -> int:
    """Worker-only artifact metadata; never extracts images or remote links."""
    reader = PdfReader(str(path))
    if reader.is_encrypted:
        raise ValueError("encrypted/password-protected PDF cannot be verified")
    pages = len(reader.pages)
    if not 0 < pages <= _PDF_PAGE_CAP:
        raise ValueError(f"PDF page verification limit exceeded ({_PDF_PAGE_CAP} pages)")
    return pages


def _extract_docx(path: Path) -> str:
    # mammoth's own Markdown writer is deprecated upstream ("generating HTML and
    # using a separate library to convert the HTML to Markdown is recommended");
    # markdownify is already the .html path's converter, so reuse it.
    with path.open("rb") as handle:
        result = mammoth.convert_to_html(handle, external_file_access=False)
    return markdownify(result.value)


def _fmt_row(cells: tuple, width: int) -> str:
    values = [("" if c is None else str(c)) for c in cells]
    values += [""] * (width - len(values))
    return "| " + " | ".join(values) + " |"


def _extract_xlsx(path: Path) -> str:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sections = []
        for name in wb.sheetnames[:_XLSX_SHEET_CAP]:
            ws = wb[name]
            rows = []
            truncated = False
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= _XLSX_ROW_CAP:
                    truncated = True
                    break
                rows.append(row)
            if not rows:
                sections.append(f"## {name}\n\n(empty sheet)")
                continue
            width = len(rows[0])
            lines = [_fmt_row(rows[0], width), "| " + " | ".join(["---"] * width) + " |"]
            lines += [_fmt_row(r, width) for r in rows[1:]]
            body = "\n".join(lines)
            if truncated:
                body += f"\n\n... [truncated at {_XLSX_ROW_CAP} rows of sheet '{name}'; more rows exist]"
            sections.append(f"## {name}\n\n{body}")
        if not sections:
            return "(empty workbook)"
        out = "\n\n".join(sections)
        if len(wb.sheetnames) > _XLSX_SHEET_CAP:
            out += _truncation_note("sheets", _XLSX_SHEET_CAP, len(wb.sheetnames))
        return out
    finally:
        wb.close()


def _extract_html(path: Path) -> str:
    html = path.read_text(encoding="utf-8", errors="replace")
    return markdownify(html)


_CONVERTERS = {
    ".docx": _extract_docx,
    ".xlsx": _extract_xlsx,
    ".html": _extract_html,
    ".htm": _extract_html,
}


def extract_text(path: Path) -> str:
    """Return markdown (or plain text) for `path`, deterministically and
    offline. `.md`/`.txt`/source-code files and anything else that decodes as
    UTF-8 text pass through unchanged -- no converter list to maintain for
    every language extension. Raises ValueError naming the format when there
    is genuinely nothing extractable: a binary format with no converter above,
    or (PDF-specific) a scanned page with no text layer at all -- or, before any
    of that, when the file is too large to read into memory at all.
    """
    size = path.stat().st_size
    if size > _MAX_BYTES:
        raise ValueError(
            f"{path.name} is {size // (1024 * 1024)}MB; refusing to read more than "
            f"{_MAX_BYTES // (1024 * 1024)}MB into memory"
        )
    ext = path.suffix.lower()
    if ext == ".pdf":
        from brain.extract_worker import run_pdf
        try:
            return run_pdf(path)
        except TimeoutError as exc:
            raise ValueError(str(exc)) from exc
    converter = _CONVERTERS.get(ext)
    if converter is not None:
        return converter(path)
    data = path.read_bytes()
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        raise ValueError(
            f"no extractable text: {ext or '(no extension)'} is not a supported text or document format"
        ) from None
